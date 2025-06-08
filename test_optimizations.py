#!/usr/bin/env python3
"""
Test script for PDF Translator optimizations

Tests the three key optimizations:
1. Paragraph placeholder system
2. XML text sanitization
3. File path consistency
"""

import os
import sys
import tempfile
import logging
from docx import Document

# Add current directory to path for imports
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from utils import prepare_text_for_translation, sanitize_for_xml, sanitize_filepath
from document_generator import document_generator

# Set up logging
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
logger = logging.getLogger(__name__)

def test_paragraph_placeholder_system():
    """Test the paragraph placeholder system"""
    logger.info("🧪 Testing paragraph placeholder system...")

    # Test text with double newlines
    test_text = "First paragraph.\n\nSecond paragraph.\n\nThird paragraph."

    # Apply placeholder system
    prepared_text = prepare_text_for_translation(test_text)

    # Check if placeholders were inserted
    expected = "First paragraph. [PARAGRAPH_BREAK] Second paragraph. [PARAGRAPH_BREAK] Third paragraph."

    if prepared_text == expected:
        logger.info("✅ Paragraph placeholder system working correctly")
        return True
    else:
        logger.error(f"❌ Paragraph placeholder failed. Expected: {expected}, Got: {prepared_text}")
        return False

def test_xml_sanitization():
    """Test XML text sanitization"""
    logger.info("🧪 Testing XML text sanitization...")

    # Test text with problematic XML characters
    test_text = "Normal text\x00with\x01control\x02characters\x03and\x08more\x0Bproblems\x0C"

    # Apply sanitization
    sanitized_text = sanitize_for_xml(test_text)

    # Check if control characters were removed
    expected = "Normal textwithcontrolcharactersandmoreproblems"

    if sanitized_text == expected:
        logger.info("✅ XML sanitization working correctly")
        return True
    else:
        logger.error(f"❌ XML sanitization failed. Expected: {expected}, Got: {sanitized_text}")
        return False

def test_filepath_sanitization():
    """Test file path sanitization"""
    logger.info("🧪 Testing file path sanitization...")

    # Test problematic file path
    test_path = r"C:\Users\Test\Documents\My<File>With:Problems|And*More?.docx"

    # Apply sanitization
    sanitized_path = sanitize_filepath(test_path)

    # Check if problematic characters were replaced
    expected = r"C:\Users\Test\Documents\My_File_With_Problems_And_More_.docx"

    if sanitized_path == expected:
        logger.info("✅ File path sanitization working correctly")
        return True
    else:
        logger.error(f"❌ File path sanitization failed. Expected: {expected}, Got: {sanitized_path}")
        return False

def test_document_generation_with_placeholders():
    """Test document generation with paragraph placeholders"""
    logger.info("🧪 Testing document generation with placeholders...")

    try:
        # Create test content with placeholders
        test_content = [
            {
                'type': 'paragraph',
                'text': 'First paragraph. [PARAGRAPH_BREAK] Second paragraph with placeholder.'
            },
            {
                'type': 'h1',
                'text': 'Test Heading with\x01control\x02characters'
            }
        ]

        # Create temporary output file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_file:
            temp_path = temp_file.name

        try:
            # Generate document
            result_path = document_generator.create_word_document_with_structure(
                test_content, temp_path, None, None
            )

            if result_path and os.path.exists(result_path):
                # Verify the document was created
                doc = Document(result_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

                logger.info(f"✅ Document generated successfully at: {result_path}")
                logger.info(f"   Document contains {len(paragraphs)} paragraphs")

                # Clean up
                os.unlink(result_path)
                return True
            else:
                logger.error("❌ Document generation failed - no file created")
                return False

        except Exception as e:
            logger.error(f"❌ Document generation failed with error: {e}")
            # Clean up on error
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            return False

    except Exception as e:
        logger.error(f"❌ Test setup failed: {e}")
        return False

def run_all_tests():
    """Run all optimization tests"""
    logger.info("🚀 Starting PDF Translator optimization tests...")

    tests = [
        test_paragraph_placeholder_system,
        test_xml_sanitization,
        test_filepath_sanitization,
        test_document_generation_with_placeholders
    ]

    passed = 0
    total = len(tests)

    for test in tests:
        try:
            if test():
                passed += 1
        except Exception as e:
            logger.error(f"❌ Test {test.__name__} failed with exception: {e}")

    logger.info(f"\n📊 Test Results: {passed}/{total} tests passed")

    if passed == total:
        logger.info("🎉 All optimizations are working correctly!")
        return True
    else:
        logger.warning(f"⚠️ {total - passed} test(s) failed. Please check the implementation.")
        return False

if __name__ == "__main__":
    success = run_all_tests()
    sys.exit(0 if success else 1)
